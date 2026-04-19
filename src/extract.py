"""
Extraction du contenu de fichiers PPTX et PDF vers un JSON structuré.
Chaque slide / page devient une entrée avec son texte et ses métadonnées.
"""

import json
from pathlib import Path
from pptx import Presentation
from pypdf import PdfReader


def extract_slide_text(slide):
    """Extrait tout le texte visible d'un slide (titres, bullets, tableaux)."""
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                texts.append(text)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        texts.append(cell_text)
    return "\n".join(texts)


def extract_slide_notes(slide):
    """Extrait les notes du présentateur."""
    if slide.has_notes_slide:
        return slide.notes_slide.notes_text_frame.text.strip()
    return ""


def extract_slide_title(slide):
    """Récupère le titre du slide s'il existe."""
    if slide.shapes.title and slide.shapes.title.text_frame.text.strip():
        return slide.shapes.title.text_frame.text.strip()
    return "(sans titre)"


def extract_pptx(filepath):
    """Extrait tout le contenu d'un fichier PPTX slide par slide."""
    prs = Presentation(filepath)
    filename = Path(filepath).name
    return [
        {
            "source_file": filename,
            "source_type": "pptx",
            "page_number": i,
            "title": extract_slide_title(slide),
            "content": extract_slide_text(slide),
            "notes": extract_slide_notes(slide),
        }
        for i, slide in enumerate(prs.slides, start=1)
    ]


def extract_pdf(filepath):
    """Extrait le texte d'un fichier PDF page par page (ignore les pages vides)."""
    reader = PdfReader(filepath)
    filename = Path(filepath).name
    pages_data = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages_data.append({
                "source_file": filename,
                "source_type": "pdf",
                "page_number": i,
                "title": f"Page {i}",
                "content": text,
                "notes": "",
            })
    return pages_data


def extract_all(data_dir="data", output_file="extracted_data.json"):
    """Parcourt tous les PPTX et PDF du dossier et sauvegarde un JSON global."""
    data_path = Path(data_dir)
    pptx_files = list(data_path.rglob("*.pptx"))
    pdf_files = list(data_path.rglob("*.pdf"))

    if not pptx_files and not pdf_files:
        print(f"Aucun fichier .pptx ou .pdf trouvé dans {data_dir}/")
        return

    print(f"{len(pptx_files)} PPTX et {len(pdf_files)} PDF trouvés.\n")

    all_items = []

    for pptx_file in pptx_files:
        print(f"[PPTX] Extraction : {pptx_file.name}")
        try:
            slides = extract_pptx(pptx_file)
            all_items.extend(slides)
            print(f"  -> {len(slides)} slides extraits")
        except Exception as e:
            print(f"  -> ERREUR : {e}")

    for pdf_file in pdf_files:
        print(f"[PDF] Extraction : {pdf_file.name}")
        try:
            pages = extract_pdf(pdf_file)
            all_items.extend(pages)
            print(f"  -> {len(pages)} pages extraites")
        except Exception as e:
            print(f"  -> ERREUR : {e}")

    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(all_items, f, ensure_ascii=False, indent=2)

    print(f"\nTotal : {len(all_items)} elements extraits")
    print(f"Sauvegarde : {output_file}")


if __name__ == "__main__":
    extract_all()
